"""Word, Excel and PowerPoint: modern (OOXML), OpenDocument and legacy binary formats."""
import re
import xml.etree.ElementTree as ET
import zipfile

from .base import Chunk, OffsetLocator, Unreadable

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
PARAS_PER_CHUNK = 500
ROWS_PER_CHUNK = 2000


def _open_zip(f):
    try:
        return zipfile.ZipFile(f)
    except zipfile.BadZipFile as e:
        raise Unreadable(f"corrupt: {e}")


# ---- docx -------------------------------------------------------------------
def _docx_paragraphs(zf, part):
    with zf.open(part) as fp:
        texts = []
        for event, el in ET.iterparse(fp, events=("end",)):
            if el.tag == W + "t" and el.text:
                texts.append(el.text)
            elif el.tag in (W + "tab", W + "br"):
                texts.append(" ")
            elif el.tag == W + "p":
                yield "".join(texts)
                texts = []
                el.clear()


def _docx_parts(zf):
    names = zf.namelist()
    main = [n for n in names if n == "word/document.xml"]
    extra = sorted(n for n in names if re.fullmatch(r"word/(header|footer)\d*\.xml|word/(footnotes|endnotes|comments)\.xml", n))
    return main + extra


def docx_chunks(f, ctx, inner):
    zf = _open_zip(f)
    for part in _docx_parts(zf):
        label = "" if part == "word/document.xml" else part.split("/")[-1].replace(".xml", "") + " "
        loc, parts, pos, count = OffsetLocator(), [], 0, 0
        for n, para in enumerate(_docx_paragraphs(zf, part), start=1):
            if not para:
                continue
            loc.add(pos, f"{label}paragraph {n}")
            parts.append(para + "\n")
            pos += len(para) + 1
            count += 1
            if count >= PARAS_PER_CHUNK:
                ctx.check_stop()
                yield Chunk("".join(parts), inner, loc)
                loc, parts, pos, count = OffsetLocator(), [], 0, 0
        if parts:
            yield Chunk("".join(parts), inner, loc)


# ---- pptx -------------------------------------------------------------------
def _slide_no(name):
    return int(re.search(r"(\d+)\.xml$", name).group(1))


def pptx_chunks(f, ctx, inner):
    zf = _open_zip(f)
    slides = sorted((n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)), key=_slide_no)
    notes = sorted((n for n in zf.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)), key=_slide_no)
    ctx.pages += len(slides)
    for part, label in [(s, "slide") for s in slides] + [(n, "notes of slide") for n in notes]:
        ctx.check_stop()
        with zf.open(part) as fp:
            text = " ".join(el.text for _, el in ET.iterparse(fp) if el.tag == A + "t" and el.text)
        if text.strip():
            where = f"{label} {_slide_no(part)}"
            yield Chunk(text, inner, lambda pos, w=where: w)


# ---- spreadsheets -------------------------------------------------------------
def cell_text(v) -> str:
    """Numbers stored as floats (IDs, account numbers) must come back as plain digits."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _is_header(values) -> bool:
    cells = [c for c in values if c]
    return bool(cells) and all(not re.fullmatch(r"[\d.,\-+/ ]+", c) for c in cells)


def sheet_chunks(sheet_name, rows, ctx, inner, col_name):
    """rows: iterable of (row_number, [cell values]). Each cell is written 'Header: value'."""
    header, loc, parts, pos, count = None, OffsetLocator(), [], 0, 0
    for r, values in rows:
        values = [cell_text(v) for v in values]
        if header is None and count == 0 and not parts and _is_header(values):
            header = values
            continue
        for i, val in enumerate(values):
            if not val:
                continue
            name = header[i] if header and i < len(header) and header[i] else col_name(i)
            seg = f"{name}: {val} | "
            loc.add(pos, f"{sheet_name}!{col_name(i)}{r}")
            parts.append(seg)
            pos += len(seg)
        parts.append("\n")
        pos += 1
        count += 1
        if count >= ROWS_PER_CHUNK:
            ctx.check_stop()
            yield Chunk("".join(parts), inner, loc)
            loc, parts, pos, count = OffsetLocator(), [], 0, 0
    if parts:
        yield Chunk("".join(parts), inner, loc)


def xlsx_chunks(f, ctx, inner):
    import openpyxl
    from openpyxl.utils import get_column_letter
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    except Exception as e:
        raise Unreadable(f"corrupt: {type(e).__name__}")
    try:
        for ws in wb.worksheets:
            rows = ((i, list(row)) for i, row in enumerate(ws.iter_rows(values_only=True), start=1))
            yield from sheet_chunks(ws.title, rows, ctx, inner, lambda i: get_column_letter(i + 1))
    finally:
        wb.close()


def xls_chunks(f, ctx, inner):
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=f.read(), on_demand=True)
    except xlrd.biffh.XLRDError as e:
        raise Unreadable("encrypted" if "encrypt" in str(e).lower() else f"corrupt: {e}")
    from openpyxl.utils import get_column_letter
    for sh in book.sheets():
        rows = ((r + 1, sh.row_values(r)) for r in range(sh.nrows))
        yield from sheet_chunks(sh.name, rows, ctx, inner, lambda i: get_column_letter(i + 1))


# ---- OpenDocument / EPUB (zip of XML/XHTML) -------------------------------------
_TAGS = re.compile(r"<[^>]+>")


def odf_chunks(f, ctx, inner):
    zf = _open_zip(f)
    parts = [n for n in zf.namelist() if n == "content.xml" or n.endswith((".xhtml", ".html", ".htm"))]
    for part in sorted(parts):
        ctx.check_stop()
        raw = zf.read(part).decode("utf-8", errors="replace")
        text = _TAGS.sub(" ", raw.replace("</text:p>", "\n").replace("</p>", "\n"))
        yield Chunk(text, inner, lambda pos, p=part: p)


# ---- legacy binary Office (OLE) ---------------------------------------------------
_UTF16_RUNS = re.compile(rb"(?:[\x20-\x7e]\x00){4,}")
_ASCII_RUNS = re.compile(rb"[\x20-\x7e\t\r\n]{6,}")


def ole_chunks(f, ctx, inner):
    """Legacy .doc/.ppt: text runs recovered from the binary streams (best effort)."""
    import olefile
    try:
        ole = olefile.OleFileIO(f)
    except Exception as e:
        raise Unreadable(f"corrupt: {type(e).__name__}")
    try:
        streams = ["/".join(s) for s in ole.listdir()]
        if "EncryptedPackage" in streams or "EncryptionInfo" in streams:
            raise Unreadable("encrypted")
        ctx.note("legacy_format_best_effort", partial=False)
        for name in streams:
            if not any(k in name for k in ("WordDocument", "PowerPoint Document", "1Table", "0Table")):
                continue
            ctx.check_stop()
            data = ole.openstream(name).read()
            runs = [m.decode("utf-16le") for m in _UTF16_RUNS.findall(data)]
            runs += [m.decode("latin-1") for m in _ASCII_RUNS.findall(data)]
            if runs:
                yield Chunk("\n".join(runs), inner, lambda pos, n=name: f"stream {n}")
    finally:
        ole.close()
