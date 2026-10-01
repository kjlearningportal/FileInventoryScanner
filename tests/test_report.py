import datetime as dt
import os

import openpyxl
import pytest

from fscan.config import build_config
from fscan.report import AGE_GROUPS, age_cutoffs, age_group, build_report
from fscan.state import State

AS_OF = dt.datetime(2026, 9, 26, 12, 0, tzinfo=dt.timezone.utc)
ROOT = r"\\fs\dept"


def ep(y, m, d):
    return dt.datetime(y, m, d, 12, 0, tzinfo=dt.timezone.utc).timestamp()


def _cfg(tmp_path):
    return build_config({"sources": [{"name": "dept", "root": ROOT}], "output_dir": str(tmp_path / "out")})


def _add(state, rel, size, modified, created=None, accessed=None, sha=None, status="done",
         group="document", ext=".docx", sensitivity=None, pii=None, topics=None):
    path = ROOT + "\\" + rel
    state.db.execute(
        "INSERT INTO files(path, source, dir, name, ext, size, created, modified, accessed, attributes, "
        "status, type_group, detected_type, sha256, seen_run) VALUES(?,?,?,?,?,?,?,?,?,0,?,?,?,?,1)",
        (path, "dept", os.path.dirname(path), os.path.basename(path), ext, size,
         created or modified, modified, accessed or modified, status, group, ext.strip("."), sha))
    if sensitivity:
        state.db.execute(
            "INSERT INTO content(path, status, content_status, sensitivity, pii_counts, pii_types, "
            "has_credentials, topics) VALUES(?, 'done', 'ok', ?, ?, ?, 0, ?)",
            (path, sensitivity, pii, (pii or "").split("=")[0], topics))


@pytest.fixture
def report(tmp_path):
    state = State(str(tmp_path / "s.db"))
    GB = 1024 ** 3
    _add(state, r"HR\Payroll\2026\march.xlsx", 2 * GB, ep(2026, 8, 1), group="spreadsheet", ext=".xlsx",
         sensitivity="HIGH", pii="AADHAAR=300", topics="HR/Payroll")                       # Active
    _add(state, r"HR\Payroll\2019\salary.xlsx", 1 * GB, ep(2019, 1, 1), group="spreadsheet", ext=".xlsx",
         sensitivity="HIGH", pii="BANK_ACCOUNT=12", topics="HR/Payroll",
         created=ep(2018, 5, 1))                                                          # Archive, stale+sensitive
    _add(state, r"HR\offer.docx", 10 * 1024 ** 2, ep(2024, 1, 1), sensitivity="MEDIUM", pii="PAN=1")  # Aging
    _add(state, r"HR\notes.docx", 1024, ep(2025, 12, 1), sensitivity="MEDIUM", pii="PAN=1")          # Recent
    _add(state, r"Finance\a.pdf", 500 * 1024 ** 2, ep(2022, 6, 1), sha="d1", group="document", ext=".pdf")  # Stale
    _add(state, r"Finance\Old\a_copy.pdf", 500 * 1024 ** 2, ep(2022, 6, 2), sha="d1", group="document", ext=".pdf")
    _add(state, r"Archive\a_copy2.pdf", 500 * 1024 ** 2, ep(2020, 1, 1), sha="d1", group="document", ext=".pdf")
    _add(state, r"empty1.txt", 0, ep(2020, 1, 1), sha="e", group="empty", ext=".txt")
    _add(state, r"empty2.txt", 0, ep(2020, 1, 1), sha="e", group="empty", ext=".txt")
    _add(state, r"tmp\~$lock.docx", 10, ep(2026, 9, 1), status="skipped")
    _add(state, r"Logs\big.log", 3 * GB, ep(2026, 9, 20), accessed=ep(2026, 9, 25), group="log", ext=".log")
    path = build_report(state, _cfg(tmp_path), as_of=AS_OF)
    wb = openpyxl.load_workbook(path, read_only=True)
    yield {name: [list(r) for r in wb[name].iter_rows(values_only=True)] for name in wb.sheetnames}
    wb.close()
    state.close()


def _table(rows):
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:]]


def test_age_group_boundaries():
    cut = age_cutoffs(AS_OF)
    names = [g[0] for g in AGE_GROUPS]
    six_months = dt.datetime(2026, 3, 26, 12, 0, tzinfo=dt.timezone.utc).timestamp()
    assert age_group(six_months, cut) == names[0]            # exactly 6 months: still Active
    assert age_group(six_months - 1, cut) == names[1]
    assert age_group(ep(2025, 9, 26), cut) == names[1]        # exactly 1 year: Recent
    assert age_group(ep(2025, 9, 26) - 1, cut) == names[2]
    assert age_group(ep(2023, 9, 26) - 1, cut) == names[3]    # just over 3 years
    assert age_group(ep(2021, 9, 26) - 1, cut) == names[4]    # just over 5 years
    assert age_group(ep(2030, 1, 1), cut) == names[0]         # future date counts as active
    assert age_group(None, cut) == "Unknown"


def test_sheets_exist(report):
    assert list(report) == ["Summary", "Age by folder", "Duplicate groups", "Duplicate files",
                            "Stale sensitive", "File types", "Largest files", "Growth by year"]


def test_age_by_folder_three_levels(report):
    rows = {r["Folder"]: r for r in _table(report["Age by folder"])}
    assert ROOT + r"\HR\Payroll\2026" in rows and ROOT + r"\HR\Payroll\2019" in rows
    assert rows[ROOT + r"\HR"]["Files"] == 2                   # offer.docx + notes.docx
    assert rows[ROOT]["Files"] == 2                            # the two empty files at the root
    payroll19 = rows[ROOT + r"\HR\Payroll\2019"]
    assert payroll19["Archive (>5y) files"] == 1 and payroll19["% stale (3y+)"] == 100.0
    assert not any("~$lock" in k for k in rows)                # skipped files excluded


def test_duplicates(report):
    groups = _table(report["Duplicate groups"])
    assert len(groups) == 1                                    # empty files never count
    g = groups[0]
    assert g["Copies"] == 3
    assert g["Wasted GB"] == pytest.approx(round(2 * 500 / 1024, 2))
    files = _table(report["Duplicate files"])
    assert {os.path.basename(r["Path"]) for r in files} == {"a.pdf", "a_copy.pdf", "a_copy2.pdf"}


def test_stale_sensitive(report):
    rows = _table(report["Stale sensitive"])
    names = [os.path.basename(r["Path"]) for r in rows]
    assert names == ["salary.xlsx", "offer.docx"]              # HIGH first; recent/active excluded
    assert rows[0]["Data found"] == "BANK_ACCOUNT=12"
    assert rows[0]["Age group"] == "Archive (>5y)"


def test_largest_and_growth(report):
    largest = _table(report["Largest files"])
    assert os.path.basename(largest[0]["Path"]) == "big.log"
    growth = {r["Year"]: r for r in _table(report["Growth by year"])}
    assert growth[2018]["Files created"] == 1                  # salary.xlsx created 2018
    assert growth[2019]["Files last modified"] == 1


def test_file_types(report):
    rows = {r["Type group"]: r for r in _table(report["File types"]) if r["Type group"]}
    assert rows["document"]["Files"] == 5
    assert rows["log"]["GB"] == pytest.approx(3.0)


def test_summary_text(report):
    text = "\n".join(" ".join(str(c) for c in r if c is not None) for r in report["Summary"])
    assert "As of 2026-09-26" in text
    assert "Excluded by config: 1" in text
    assert "last-accessed" in text.lower()
    assert "Wasted" in text


def test_access_time_check(tmp_path):
    from fscan.report import access_time_recorded
    assert access_time_recorded(n_files=1000, n_accessed_after_modified=5) is False
    assert access_time_recorded(n_files=1000, n_accessed_after_modified=300) is True


def test_row_limit_truncates_and_says_so(tmp_path):
    state = State(str(tmp_path / "s.db"))
    for i in range(30):
        _add(state, rf"X\f{i}.docx", 100 + i, ep(2019, 1, 1), sensitivity="HIGH", pii="PAN=20")
    path = build_report(state, _cfg(tmp_path), as_of=AS_OF, max_rows=10)
    wb = openpyxl.load_workbook(path, read_only=True)
    assert sum(1 for _ in wb["Stale sensitive"].iter_rows()) == 11   # header + 10
    summary = " ".join(str(c) for r in wb["Summary"].iter_rows(values_only=True) for c in r if c)
    assert "truncated" in summary.lower()
    wb.close()
    state.close()
